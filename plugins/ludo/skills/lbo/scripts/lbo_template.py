#!/usr/bin/env python3
"""Build a live, formula-driven LBO (leveraged buyout) workbook.

Usage:
  python3 lbo_template.py lbo.xlsx                  # sample inputs, clearly marked
  python3 lbo_template.py lbo.xlsx --inputs a.json  # your own inputs (same keys as SAMPLE)

Needs: pip install openpyxl
Blue cells are inputs. Everything else is a formula: sources and uses, a 5-year
operating model, a debt schedule with a cash sweep, and returns (IRR and MOIC).
"""
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is missing. Run: pip install openpyxl")

SAMPLE = {
    "company": "Sample Co. (made-up numbers, replace with real ones)",
    "ltm_revenue": 500.0,
    "ltm_ebitda": 100.0,
    "entry_multiple": 10.0,
    "debt_multiple": 5.0,
    "interest_rate": 0.08,
    "fees_pct": 0.02,
    "min_cash": 20.0,
    "growth": [0.08, 0.08, 0.07, 0.06, 0.05],
    "ebitda_margin": [0.20, 0.21, 0.21, 0.22, 0.22],
    "da_pct": 0.03,
    "capex_pct": 0.035,
    "nwc_pct": 0.10,
    "tax_rate": 0.25,
    "exit_multiple": 10.0,
}

BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
INP = PatternFill("solid", fgColor="EEF3FB")
HEAD = PatternFill("solid", fgColor="F2F2F2")
NUM = '#,##0.0;(#,##0.0)'
PCT = "0.0%"
MULT = '0.0"x"'


def build(v, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "LBO"
    ws["A1"] = f"{v['company']}: leveraged buyout"
    ws["A1"].font = TITLE
    ws["A2"] = "Units: $ millions. Blue cells are inputs; everything else updates on its own."

    # Inputs block
    inputs = [
        ("LTM revenue", "ltm_revenue", NUM), ("LTM EBITDA", "ltm_ebitda", NUM),
        ("Entry EV / EBITDA", "entry_multiple", MULT), ("Debt / EBITDA at entry", "debt_multiple", MULT),
        ("Interest rate on debt", "interest_rate", PCT), ("Transaction fees (% of EV)", "fees_pct", PCT),
        ("Minimum cash kept", "min_cash", NUM), ("D&A (% of revenue)", "da_pct", PCT),
        ("Capex (% of revenue)", "capex_pct", PCT), ("Net working capital (% of revenue)", "nwc_pct", PCT),
        ("Tax rate", "tax_rate", PCT), ("Exit EV / EBITDA", "exit_multiple", MULT),
    ]
    ref = {}
    ws["A4"] = "Inputs"
    ws["A4"].font = BOLD
    for i, (lab, key, fmt) in enumerate(inputs, start=5):
        ws.cell(i, 1, lab)
        c = ws.cell(i, 2, v[key])
        c.font, c.fill, c.number_format = BLUE, INP, fmt
        ref[key] = f"$B${i}"

    # Sources and uses
    ws["D4"] = "Sources and uses"
    ws["D4"].font = BOLD
    su = [
        ("Purchase enterprise value", f"={ref['ltm_ebitda']}*{ref['entry_multiple']}"),
        ("Transaction fees", f"=E5*{ref['fees_pct']}"),
        ("Minimum cash funded", f"={ref['min_cash']}"),
        ("Total uses", "=E5+E6+E7"),
        ("Debt raised", f"={ref['ltm_ebitda']}*{ref['debt_multiple']}"),
        ("Sponsor equity", "=E8-E9"),
        ("Equity % of total", "=E10/E8"),
    ]
    for i, (lab, f) in enumerate(su, start=5):
        ws.cell(i, 4, lab)
        c = ws.cell(i, 5, f)
        c.number_format = PCT if "%" in lab else NUM
    for r in (8, 10):
        ws.cell(r, 4).font = BOLD
        ws.cell(r, 5).font = BOLD

    # Operating model and debt schedule: columns C (year 0) to H (year 5)
    top = 19
    ws.cell(top, 1, "Year").font = BOLD
    for j in range(6):
        c = ws.cell(top, 3 + j, j)
        c.font, c.fill = BOLD, HEAD
    rows = ["Revenue growth", "Revenue", "EBITDA margin", "EBITDA", "D&A", "EBIT", "Interest", "Pre-tax income",
            "Taxes", "Net income", "+ D&A", "- Capex", "- Increase in NWC", "Free cash flow for debt paydown",
            "Debt, start of year", "Debt repaid", "Debt, end of year", "Debt / EBITDA", "Cash, end of year"]
    R = {name: top + 1 + i for i, name in enumerate(rows)}
    for name, r in R.items():
        ws.cell(r, 1, name)
    for key in ("Net income", "Free cash flow for debt paydown", "Debt, end of year"):
        ws.cell(R[key], 1).font = BOLD

    ws.cell(R["Revenue"], 3, f"={ref['ltm_revenue']}")
    ws.cell(R["EBITDA"], 3, f"={ref['ltm_ebitda']}")
    ws.cell(R["Debt, end of year"], 3, "=E9")
    ws.cell(R["Cash, end of year"], 3, f"={ref['min_cash']}")
    ws.cell(R["Debt / EBITDA"], 3, f"=C{R['Debt, end of year']}/C{R['EBITDA']}")
    for j in range(1, 6):
        col = 3 + j
        C = get_column_letter(col)
        P = get_column_letter(col - 1)
        g = ws.cell(R["Revenue growth"], col, v["growth"][j - 1])
        m = ws.cell(R["EBITDA margin"], col, v["ebitda_margin"][j - 1])
        for c in (g, m):
            c.font, c.fill = BLUE, INP
        f = {
            "Revenue": f"={P}{R['Revenue']}*(1+{C}{R['Revenue growth']})",
            "EBITDA": f"={C}{R['Revenue']}*{C}{R['EBITDA margin']}",
            "D&A": f"=-{C}{R['Revenue']}*{ref['da_pct']}",
            "EBIT": f"={C}{R['EBITDA']}+{C}{R['D&A']}",
            "Interest": f"=-{P}{R['Debt, end of year']}*{ref['interest_rate']}",
            "Pre-tax income": f"={C}{R['EBIT']}+{C}{R['Interest']}",
            "Taxes": f"=-MAX(0,{C}{R['Pre-tax income']})*{ref['tax_rate']}",
            "Net income": f"={C}{R['Pre-tax income']}+{C}{R['Taxes']}",
            "+ D&A": f"=-{C}{R['D&A']}",
            "- Capex": f"=-{C}{R['Revenue']}*{ref['capex_pct']}",
            "- Increase in NWC": f"=-({C}{R['Revenue']}-{P}{R['Revenue']})*{ref['nwc_pct']}",
            "Free cash flow for debt paydown": f"={C}{R['Net income']}+{C}{R['+ D&A']}+{C}{R['- Capex']}+{C}{R['- Increase in NWC']}",
            "Debt, start of year": f"={P}{R['Debt, end of year']}",
            "Debt repaid": f"=MIN({C}{R['Debt, start of year']},MAX(0,{C}{R['Free cash flow for debt paydown']}))",
            "Debt, end of year": f"={C}{R['Debt, start of year']}-{C}{R['Debt repaid']}",
            "Debt / EBITDA": f"={C}{R['Debt, end of year']}/{C}{R['EBITDA']}",
            "Cash, end of year": f"={P}{R['Cash, end of year']}+{C}{R['Free cash flow for debt paydown']}-{C}{R['Debt repaid']}",
        }
        for name, formula in f.items():
            ws.cell(R[name], col, formula)
    for name, r in R.items():
        for col in range(3, 9):
            ws.cell(r, col).number_format = PCT if name in ("Revenue growth", "EBITDA margin") else (
                MULT if name == "Debt / EBITDA" else NUM)

    # Returns
    rr = R["Cash, end of year"] + 2
    ws.cell(rr, 1, "Returns").font = BOLD
    ret = [
        ("Exit EBITDA (year 5)", f"=H{R['EBITDA']}", NUM),
        ("Exit enterprise value", f"=B{rr + 1}*{ref['exit_multiple']}", NUM),
        ("Less: debt at exit", f"=-H{R['Debt, end of year']}", NUM),
        ("Plus: cash at exit", f"=H{R['Cash, end of year']}", NUM),
        ("Exit equity value", f"=B{rr + 2}+B{rr + 3}+B{rr + 4}", NUM),
        ("Sponsor equity invested", "=E10", NUM),
        ("MOIC (money multiple)", f"=B{rr + 5}/B{rr + 6}", MULT),
        ("IRR", f"=(B{rr + 7})^(1/5)-1", PCT),
    ]
    for i, (lab, f, fmt) in enumerate(ret, start=rr + 1):
        ws.cell(i, 1, lab)
        c = ws.cell(i, 2, f)
        c.number_format = fmt
    for i in (rr + 7, rr + 8):
        ws.cell(i, 1).font = BOLD
        ws.cell(i, 2).font = BOLD
    ws.cell(rr + 10, 1, "Returns bridge: how much came from EBITDA growth, multiple change and debt paydown").font = Font(italic=True)
    ws.cell(rr + 11, 1, "EBITDA growth")
    ws.cell(rr + 11, 2, f"=(B{rr + 1}-{ref['ltm_ebitda']})*{ref['entry_multiple']}").number_format = NUM
    ws.cell(rr + 12, 1, "Multiple change")
    ws.cell(rr + 12, 2, f"=({ref['exit_multiple']}-{ref['entry_multiple']})*B{rr + 1}").number_format = NUM
    ws.cell(rr + 13, 1, "Debt paydown and cash build")
    ws.cell(rr + 13, 2, f"=(E9-H{R['Debt, end of year']})+(H{R['Cash, end of year']}-{ref['min_cash']})").number_format = NUM
    ws.cell(rr + 14, 1, "Less: fees")
    ws.cell(rr + 14, 2, "=-E6").number_format = NUM
    ws.cell(rr + 15, 1, "Check: bridge equals equity gain")
    ws.cell(rr + 15, 2, f'=IF(ABS(SUM(B{rr + 11}:B{rr + 14})-(B{rr + 5}-B{rr + 6}))<0.01,"OK","ERROR")')

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["D"].width = 26
    for col in "CEFGH":
        ws.column_dimensions[col].width = 12
    wb.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    vals = dict(SAMPLE)
    if "--inputs" in sys.argv:
        vals.update(json.load(open(sys.argv[sys.argv.index("--inputs") + 1])))
    build(vals, sys.argv[1])
