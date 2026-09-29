#!/usr/bin/env python3
"""Build a linked three-statement model: income statement, balance sheet and cash flow.

Usage:
  python3 three_statement.py model.xlsx                  # sample inputs, clearly marked
  python3 three_statement.py model.xlsx --inputs a.json  # your own inputs (same keys as SAMPLE)

Needs: pip install openpyxl
Blue cells are inputs. Every other cell is a formula. Cash comes from the cash flow
statement, and a check row proves assets = liabilities + equity every year.
"""
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is missing. Run: pip install openpyxl")

SAMPLE = {
    "company": "Sample Co. (made-up numbers, replace with real ones)",
    "base_year": 2025,
    "opening": {"revenue": 400.0, "cash": 50.0, "receivables": 45.0, "inventory": 60.0, "ppe": 200.0,
                "payables": 30.0, "debt": 150.0, "equity": 175.0},
    "growth": [0.10, 0.09, 0.08, 0.07, 0.06],
    "gross_margin": 0.42,
    "sga_pct": 0.22,
    "depreciation_pct_ppe": 0.10,
    "capex_pct": 0.06,
    "receivable_days": 40,
    "inventory_days": 90,
    "payable_days": 45,
    "interest_rate": 0.07,
    "tax_rate": 0.25,
    "debt_repayment": 20.0,
    "dividend_payout": 0.20,
}

BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
INP = PatternFill("solid", fgColor="EEF3FB")
HEAD = PatternFill("solid", fgColor="F2F2F2")
NUM = '#,##0.0;(#,##0.0)'
PCT = "0.0%"


def build(v, out):
    wb = Workbook()
    a = wb.active
    a.title = "Inputs"
    a["A1"] = f"{v['company']}: three-statement model inputs"
    a["A1"].font = TITLE
    a["A2"] = "Units: $ millions. Blue cells are inputs."
    scalars = [("Gross margin", "gross_margin", PCT), ("SG&A (% of revenue)", "sga_pct", PCT),
               ("Depreciation (% of opening PP&E)", "depreciation_pct_ppe", PCT), ("Capex (% of revenue)", "capex_pct", PCT),
               ("Receivable days", "receivable_days", "0"), ("Inventory days (of COGS)", "inventory_days", "0"),
               ("Payable days (of COGS)", "payable_days", "0"), ("Interest rate on opening debt", "interest_rate", PCT),
               ("Tax rate", "tax_rate", PCT), ("Debt repaid each year", "debt_repayment", NUM),
               ("Dividend payout (% of net income)", "dividend_payout", PCT)]
    ref = {}
    for i, (lab, key, fmt) in enumerate(scalars, start=4):
        a.cell(i, 1, lab)
        c = a.cell(i, 2, v[key])
        c.font, c.fill, c.number_format = BLUE, INP, fmt
        ref[key] = f"Inputs!$B${i}"
    r0 = 4 + len(scalars) + 1
    a.cell(r0, 1, "Opening balance sheet and revenue (base year)").font = BOLD
    for i, (lab, key) in enumerate([("Revenue", "revenue"), ("Cash", "cash"), ("Receivables", "receivables"),
                                     ("Inventory", "inventory"), ("PP&E, net", "ppe"), ("Payables", "payables"),
                                     ("Debt", "debt"), ("Equity", "equity")], start=r0 + 1):
        a.cell(i, 1, lab)
        c = a.cell(i, 2, v["opening"][key])
        c.font, c.fill, c.number_format = BLUE, INP, NUM
        ref["o_" + key] = f"Inputs!$B${i}"
    rg = r0 + 10
    a.cell(rg, 1, "Revenue growth").font = BOLD
    for j in range(5):
        a.cell(rg - 1, 2 + j, v["base_year"] + 1 + j).font = BOLD
        c = a.cell(rg, 2 + j, v["growth"][j])
        c.font, c.fill, c.number_format = BLUE, INP, PCT
    a.column_dimensions["A"].width = 38

    m = wb.create_sheet("Model")
    m["A1"] = f"{v['company']}: linked three-statement model"
    m["A1"].font = TITLE
    years = [v["base_year"] + j for j in range(6)]
    for j, y in enumerate(years):
        c = m.cell(3, 2 + j, f"{y}{'A' if j == 0 else 'E'}")
        c.font, c.fill = BOLD, HEAD
    lines = [
        ("INCOME STATEMENT", None), ("Revenue", "rev"), ("Cost of goods sold", "cogs"), ("Gross profit", "gp"),
        ("SG&A", "sga"), ("Depreciation", "dep"), ("Operating income (EBIT)", "ebit"), ("Interest", "int"),
        ("Pre-tax income", "pti"), ("Taxes", "tax"), ("Net income", "ni"), ("", None),
        ("BALANCE SHEET", None), ("Cash", "cash"), ("Receivables", "ar"), ("Inventory", "inv"), ("PP&E, net", "ppe"),
        ("Total assets", "ta"), ("Payables", "ap"), ("Debt", "debt"), ("Equity", "eq"),
        ("Total liabilities and equity", "tle"), ("Check: balances", "chk"), ("", None),
        ("CASH FLOW STATEMENT", None), ("Net income", "cf_ni"), ("+ Depreciation", "cf_dep"),
        ("- Increase in receivables", "cf_ar"), ("- Increase in inventory", "cf_inv"), ("+ Increase in payables", "cf_ap"),
        ("Cash from operations", "cfo"), ("- Capex", "capex"), ("Cash from investing", "cfi"),
        ("- Debt repaid", "cf_debt"), ("- Dividends", "div"), ("Cash from financing", "cff"),
        ("Net change in cash", "dcash"),
    ]
    R = {}
    for i, (lab, key) in enumerate(lines, start=4):
        cell = m.cell(i, 1, lab)
        if key is None and lab:
            cell.font = BOLD
        if key in ("gp", "ebit", "ni", "ta", "tle", "cfo", "dcash"):
            cell.font = BOLD
        if key:
            R[key] = i

    # Base year column B: balance sheet from inputs, revenue only on the income statement
    m.cell(R["rev"], 2, f"={ref['o_revenue']}")
    for key, o in (("cash", "cash"), ("ar", "receivables"), ("inv", "inventory"), ("ppe", "ppe"),
                   ("ap", "payables"), ("debt", "debt"), ("eq", "equity")):
        m.cell(R[key], 2, f"={ref['o_' + o]}")
    m.cell(R["ta"], 2, f"=B{R['cash']}+B{R['ar']}+B{R['inv']}+B{R['ppe']}")
    m.cell(R["tle"], 2, f"=B{R['ap']}+B{R['debt']}+B{R['eq']}")
    m.cell(R["chk"], 2, f'=IF(ABS(B{R["ta"]}-B{R["tle"]})<0.01,"OK","ERROR")')

    for j in range(1, 6):
        C = get_column_letter(2 + j)
        P = get_column_letter(1 + j)
        g = f"Inputs!{get_column_letter(1 + j)}{rg}"
        f = {
            "rev": f"={P}{R['rev']}*(1+{g})",
            "cogs": f"=-{C}{R['rev']}*(1-{ref['gross_margin']})",
            "gp": f"={C}{R['rev']}+{C}{R['cogs']}",
            "sga": f"=-{C}{R['rev']}*{ref['sga_pct']}",
            "dep": f"=-{P}{R['ppe']}*{ref['depreciation_pct_ppe']}",
            "ebit": f"={C}{R['gp']}+{C}{R['sga']}+{C}{R['dep']}",
            "int": f"=-{P}{R['debt']}*{ref['interest_rate']}",
            "pti": f"={C}{R['ebit']}+{C}{R['int']}",
            "tax": f"=-MAX(0,{C}{R['pti']})*{ref['tax_rate']}",
            "ni": f"={C}{R['pti']}+{C}{R['tax']}",
            "ar": f"={C}{R['rev']}*{ref['receivable_days']}/365",
            "inv": f"=-{C}{R['cogs']}*{ref['inventory_days']}/365",
            "ppe": f"={P}{R['ppe']}+{C}{R['dep']}-{C}{R['capex']}",
            "ap": f"=-{C}{R['cogs']}*{ref['payable_days']}/365",
            "debt": f"=MAX(0,{P}{R['debt']}-{ref['debt_repayment']})",
            "eq": f"={P}{R['eq']}+{C}{R['ni']}+{C}{R['div']}",
            "cash": f"={P}{R['cash']}+{C}{R['dcash']}",
            "ta": f"={C}{R['cash']}+{C}{R['ar']}+{C}{R['inv']}+{C}{R['ppe']}",
            "tle": f"={C}{R['ap']}+{C}{R['debt']}+{C}{R['eq']}",
            "chk": f'=IF(ABS({C}{R["ta"]}-{C}{R["tle"]})<0.01,"OK","ERROR")',
            "cf_ni": f"={C}{R['ni']}",
            "cf_dep": f"=-{C}{R['dep']}",
            "cf_ar": f"=-({C}{R['ar']}-{P}{R['ar']})",
            "cf_inv": f"=-({C}{R['inv']}-{P}{R['inv']})",
            "cf_ap": f"={C}{R['ap']}-{P}{R['ap']}",
            "cfo": f"=SUM({C}{R['cf_ni']}:{C}{R['cf_ap']})",
            "capex": f"=-{C}{R['rev']}*{ref['capex_pct']}",
            "cfi": f"={C}{R['capex']}",
            "cf_debt": f"=-({P}{R['debt']}-{C}{R['debt']})",
            "div": f"=-MAX(0,{C}{R['ni']})*{ref['dividend_payout']}",
            "cff": f"={C}{R['cf_debt']}+{C}{R['div']}",
            "dcash": f"={C}{R['cfo']}+{C}{R['cfi']}+{C}{R['cff']}",
        }
        for key, formula in f.items():
            m.cell(R[key], 2 + j, formula)
    for key, r in R.items():
        if key != "chk":
            for col in range(2, 8):
                m.cell(r, col).number_format = NUM
    m.column_dimensions["A"].width = 32
    for j in range(6):
        m.column_dimensions[get_column_letter(2 + j)].width = 12
    wb.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    vals = dict(SAMPLE)
    if "--inputs" in sys.argv:
        vals.update(json.load(open(sys.argv[sys.argv.index("--inputs") + 1])))
    build(vals, sys.argv[1])
