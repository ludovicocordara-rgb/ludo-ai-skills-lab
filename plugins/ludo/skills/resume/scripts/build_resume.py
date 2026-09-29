#!/usr/bin/env python3
"""Build a clean one-page resume (.docx) from a JSON file.

Usage:
  python3 build_resume.py resume.json Lastname_Firstname_Resume.docx
  python3 build_resume.py resume.json Lastname_Firstname_Resume.docx --pdf   # also makes a PDF (needs LibreOffice)

Needs: pip install python-docx
See ../examples/sample.json for the JSON shape.
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    sys.exit("python-docx is missing. Run: pip install python-docx")

FONT = "Times New Roman"
BODY = 10.5
MARGIN = 0.75
TEXT_WIDTH = 8.5 - 2 * MARGIN


def style_run(run, size=BODY, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def tight(p, before=0, after=0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.0
    return p


def add_link(p, text, url):
    """Blue underlined hyperlink inside paragraph p."""
    rid = p.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                           is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(a), FONT)
    rpr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rpr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(BODY * 2)))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    link.append(r)
    p._p.append(link)


def rule_below(p):
    """Thin black line under a section header."""
    ppr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    bdr.append(bottom)
    ppr.append(bdr)


def bullet(doc, runs):
    p = tight(doc.add_paragraph())
    pf = p.paragraph_format
    pf.left_indent = Inches(0.25)
    pf.first_line_indent = Inches(-0.15)
    style_run(p.add_run("•  "))
    for text, fmt in runs:
        style_run(p.add_run(text), **fmt)


def build(spec, out):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.6)
        s.left_margin = s.right_margin = Inches(MARGIN)

    name = tight(doc.add_paragraph(), after=2)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(name.add_run(spec["name"]), size=14, bold=True)

    c = spec.get("contact", {})
    line = tight(doc.add_paragraph(), after=4)
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = []
    if c.get("city"):
        parts.append(("text", c["city"]))
    if c.get("phone"):
        parts.append(("text", c["phone"]))
    if c.get("email"):
        parts.append(("link", c["email"], "mailto:" + c["email"]))
    if c.get("linkedin"):
        url = c["linkedin"] if c["linkedin"].startswith("http") else "https://" + c["linkedin"]
        parts.append(("link", c["linkedin"], url))
    for i, part in enumerate(parts):
        if i:
            style_run(line.add_run("  |  "))
        if part[0] == "text":
            style_run(line.add_run(part[1]))
        else:
            add_link(line, part[1], part[2])

    for section in spec["sections"]:
        h = tight(doc.add_paragraph(), before=6, after=3)
        style_run(h.add_run(section["header"]), size=11, bold=True)
        rule_below(h)

        for e in section.get("entries", []):
            p = tight(doc.add_paragraph(), before=2)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(TEXT_WIDTH), WD_TAB_ALIGNMENT.RIGHT)
            style_run(p.add_run(e["org"]), bold=True, italic=True)
            if e.get("loc"):
                style_run(p.add_run(e["loc"]), italic=True)
            if e.get("date"):
                style_run(p.add_run("\t" + e["date"]))
            if e.get("title"):
                t = tight(doc.add_paragraph())
                style_run(t.add_run(e["title"]), italic=True)
            for label, rest in e.get("labelled", []):
                bullet(doc, [(label, {"bold": True, "italic": True}), (rest, {})])
            for b in e.get("bullets", []):
                bullet(doc, [(b, {})])

        for label, rest in section.get("kv", []):
            p = tight(doc.add_paragraph())
            style_run(p.add_run(label), bold=True)
            style_run(p.add_run(rest))

    doc.save(out)
    return out


def to_pdf(docx_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice not found, so no PDF. Open the .docx in Word or Google Docs and export to PDF instead.")
        return None
    outdir = str(Path(docx_path).resolve().parent)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", outdir, docx_path],
                   check=True, capture_output=True)
    pdf = str(Path(docx_path).with_suffix(".pdf"))
    try:
        import fitz  # PyMuPDF, optional
        pages = fitz.open(pdf).page_count
        print(f"PDF has {pages} page(s).")
        if pages != 1:
            print("Over one page. Shorten the longest bullets first, never shrink the font.")
            sys.exit(2)
    except ImportError:
        print("PDF written. Open it and check it fits on one page.")
    return pdf


if __name__ == "__main__":
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    spec = json.loads(Path(sys.argv[1]).read_text())
    out = build(spec, sys.argv[2])
    print("Wrote", out)
    if "--pdf" in sys.argv:
        to_pdf(out)
