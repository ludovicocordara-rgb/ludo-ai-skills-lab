#!/usr/bin/env python3
"""Build a live comparable-companies (comps) workbook.

Usage:
  python3 comps_template.py comps.xlsx                    # sample peers, clearly marked, to learn the layout
  python3 comps_template.py comps.xlsx --inputs data.json # your own companies (same keys as SAMPLE)

Needs: pip install openpyxl
Inputs tab: raw numbers with a source for each row (blue = input).
Comps tab: every multiple and statistic is a formula. Negative or missing values show "NM".
"""
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is missing. Run: pip install openpyxl")

SAMPLE = {
    "title": "Sample comps (made-up numbers, replace with real ones)",
    "units": "$ millions, except per share",
    "target": "Target Co",
    "companies": [
        # name, ticker, price, diluted shares, debt, cash, LTM revenue, LTM EBITDA, LTM net income, revenue growth, source
        ["Target Co", "TGT0", 42.0, 50.0, 300.0, 120.0, 960.0, 163.0, 72.0, 0.09, "10-K FY2025; price as of <date>"],
        ["Peer A", "PEA", 65.0, 80.0, 900.0, 200.0, 2400.0, 430.0, 210.0, 0.06, "10-K FY2025; price as of <date>"],
        ["Peer B", "PEB", 18.0, 120.0, 150.0, 90.0, 1100.0, 150.0, 60.0, 0.12, "10-K FY2025; price as of <date>"],
        ["Peer C", "PEC", 110.0, 30.0, 400.0, 50.0, 1500.0, 260.0, 140.0, 0.04, "10-K FY2025; price as of <date>"],
        ["Peer D", "PED", 9.5, 200.0, 50.0, 300.0, 700.0, -20.0, -45.0, 0.25, "10-K FY2025; price as of <date>"],
        ["Peer E", "PEE", 55.0, 45.0, 600.0, 80.0, 1300.0, 210.0, 95.0, 0.07, "10-K FY2025; price as of <date>"],
    ],
}

BLUE = Font(color="1F4E9E")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEAD = PatternFill("solid", fgColor="F2F2F2")
TARGET = PatternFill("solid", fgColor="FFF4E5")
NUM = '#,##0.0;(#,##0.0)'
MULT = '0.0"x"'
PCT = "0.0%"


def nm(expr, denom):
    """Multiple that shows NM when the denominator is missing or not positive."""
    return f'=IF(OR({denom}="",{denom}<=0),"NM",{expr})'


def build(v, out):
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = v["title"]
    ws["A1"].font = TITLE
    ws["A2"] = f"Units: {v['units']}. Blue cells are inputs. Put the source and date of every row in column K."
    heads = ["Company", "Ticker", "Share price", "Diluted shares", "Total debt", "Cash", "LTM revenue",
             "LTM EBITDA", "LTM net income", "Revenue growth", "Source"]
    for j, h in enumerate(heads, start=1):
        c = ws.cell(4, j, h)
        c.font, c.fill = BOLD, HEAD
    for i, row in enumerate(v["companies"], start=5):
        for j, val in enumerate(row, start=1):
            c = ws.cell(i, j, val)
            if 3 <= j <= 10:
                c.font = BLUE
                c.number_format = PCT if j == 10 else ("0.00" if j == 3 else NUM)
    n = len(v["companies"])
    for col, w in zip("ABCDEFGHIJK", [18, 8, 11, 13, 11, 11, 13, 12, 14, 14, 36]):
        ws.column_dimensions[col].width = w

    c = wb.create_sheet("Comps")
    c["A1"] = "Trading comparables"
    c["A1"].font = TITLE
    c["A2"] = "All figures are formulas from the Inputs tab. NM = not meaningful (negative or missing)."
    cols = ["Company", "Market cap", "Enterprise value", "EV / Revenue", "EV / EBITDA", "P / E",
            "Revenue growth", "EBITDA margin"]
    for j, h in enumerate(cols, start=1):
        cell = c.cell(4, j, h)
        cell.font, cell.fill = BOLD, HEAD
    for i in range(n):
        r = src = 5 + i
        c.cell(r, 1, f"=Inputs!A{src}")
        c.cell(r, 2, f"=Inputs!C{src}*Inputs!D{src}").number_format = NUM
        c.cell(r, 3, f"=B{r}+Inputs!E{src}-Inputs!F{src}").number_format = NUM
        c.cell(r, 4, nm(f"C{r}/Inputs!G{src}", f"Inputs!G{src}")).number_format = MULT
        c.cell(r, 5, nm(f"C{r}/Inputs!H{src}", f"Inputs!H{src}")).number_format = MULT
        c.cell(r, 6, nm(f"B{r}/Inputs!I{src}", f"Inputs!I{src}")).number_format = MULT
        c.cell(r, 7, f"=Inputs!J{src}").number_format = PCT
        c.cell(r, 8, f'=IF(Inputs!G{src}>0,Inputs!H{src}/Inputs!G{src},"NM")').number_format = PCT
        if v["companies"][i][0] == v["target"]:
            for j in range(1, 9):
                c.cell(r, j).fill = TARGET

    # Peer statistics exclude the target row
    peer_rows = [5 + i for i in range(n) if v["companies"][i][0] != v["target"]]
    stats = [("Maximum", "MAX"), ("75th percentile", "QUARTILE3"), ("Median", "MEDIAN"), ("Mean", "AVERAGE"),
             ("25th percentile", "QUARTILE1"), ("Minimum", "MIN")]
    start = 6 + n
    c.cell(start - 1, 1, "Peer statistics (target excluded)").font = BOLD
    for k, (label, fn) in enumerate(stats):
        r = start + k
        c.cell(r, 1, label).font = BOLD if label == "Median" else Font()
        for j in range(4, 9):
            L = get_column_letter(j)
            rng = ",".join(f"{L}{pr}" for pr in peer_rows)
            if fn == "QUARTILE3":
                f = f"=QUARTILE(({rng}),3)"
            elif fn == "QUARTILE1":
                f = f"=QUARTILE(({rng}),1)"
            else:
                f = f"={fn}({rng})"
            cell = c.cell(r, j, f)
            cell.number_format = PCT if j >= 7 else MULT
            if label == "Median":
                cell.font = BOLD
    c.cell(start + len(stats) + 1, 1,
           "Lead with the median: one unusual peer cannot drag it the way it drags the mean.").font = Font(italic=True)
    c.column_dimensions["A"].width = 30
    for col in "BCDEFGH":
        c.column_dimensions[col].width = 15
    wb.save(out)
    print("Wrote", out)


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    vals = dict(SAMPLE)
    if "--inputs" in sys.argv:
        vals.update(json.load(open(sys.argv[sys.argv.index("--inputs") + 1])))
    build(vals, sys.argv[1])
